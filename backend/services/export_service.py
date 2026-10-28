# export_service.py
# Builds an Excel workbook containing the records accessible to the logged-in user.

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import Config


EXPORT_DIR = Path(__file__).resolve().parent.parent / "exports"


def safe(value):
    """
    Prevent Excel formula injection.

    Values beginning with =, +, -, or @ are treated as text.
    """
    if value is None:
        return ""

    value = str(value)

    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value

    return value


def add_sheet(workbook, title, headers, rows):
    """
    Add a formatted worksheet to the workbook.
    """
    ws = workbook.create_sheet(title)

    # Header row
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=safe(header))
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Data rows
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(
                row=row_index,
                column=column_index,
                value=safe(value),
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # Freeze header
    ws.freeze_panes = "A2"

    # Add autofilter
    if headers:
        last_column = get_column_letter(len(headers))
        last_row = max(ws.max_row, 1)
        ws.auto_filter.ref = f"A1:{last_column}{last_row}"

    # Adjust column widths
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)

        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)

            for line in value.split("\n"):
                max_length = max(max_length, len(line))

        width = min(max(max_length + 2, 10), 45)

        ws.column_dimensions[column_letter].width = width

    return ws


def build_workbook(conn, user):
    """
    Build an Excel workbook containing all records accessible
    to the current user.

    Doctors:
        Can access all patients.

    Clinicians:
        Can access only patients they created.
    """

    role = user["role"]
    user_db_id = user["db_id"]

    # ---------------------------------------------------------
    # Access scope
    # ---------------------------------------------------------

    if role == "clinician":
        scope = "WHERE p.created_by = ?"
        params = (user_db_id,)
    else:
        scope = ""
        params = ()

    wb = Workbook()

    # Remove the default worksheet.
    default_sheet = wb.active
    wb.remove(default_sheet)

    # ---------------------------------------------------------
    # Patients
    # ---------------------------------------------------------

    rows = conn.execute(
        "SELECT p.patient_code, p.age, p.sex, "
        "p.created_at, u.user_id "
        "FROM patients p "
        "JOIN users u ON p.created_by = u.id "
        f"{scope} "
        "ORDER BY p.created_at DESC, p.id DESC",
        params,
    ).fetchall()

    add_sheet(
        wb,
        "Patients",
        [
            "Patient ID",
            "Age",
            "Sex",
            "Created (UTC)",
            "Created By",
        ],
        [tuple(row) for row in rows],
    )

    # ---------------------------------------------------------
    # Vitals
    # ---------------------------------------------------------

    rows = conn.execute(
        "SELECT p.patient_code, v.recorded_at, "
        "v.heart_rate, v.systolic_bp, v.diastolic_bp, "
        "v.spo2, v.temperature_c, v.respiratory_rate, "
        "u.user_id "
        "FROM vitals v "
        "JOIN patients p ON v.patient_id = p.id "
        "JOIN users u ON v.recorded_by = u.id "
        f"{scope} "
        "ORDER BY v.recorded_at DESC, v.id DESC",
        params,
    ).fetchall()

    add_sheet(
        wb,
        "Vitals",
        [
            "Patient ID",
            "Time (UTC)",
            "Heart Rate",
            "Systolic BP",
            "Diastolic BP",
            "SpO2",
            "Temperature",
            "Respiratory Rate",
            "Recorded By",
        ],
        [tuple(row) for row in rows],
    )

    # ---------------------------------------------------------
    # Clinical
    # ---------------------------------------------------------

    rows = conn.execute(
        "SELECT p.patient_code, c.recorded_at, "
        "c.symptoms, c.symptom_onset, c.medical_history, "
        "c.existing_conditions, c.clinical_observations, "
        "u.user_id "
        "FROM clinical_records c "
        "JOIN patients p ON c.patient_id = p.id "
        "JOIN users u ON c.recorded_by = u.id "
        f"{scope} "
        "ORDER BY c.recorded_at DESC, c.id DESC",
        params,
    ).fetchall()

    add_sheet(
        wb,
        "Clinical",
        [
            "Patient ID",
            "Time (UTC)",
            "Symptoms",
            "Symptom Onset",
            "Medical History",
            "Existing Conditions",
            "Clinical Observations",
            "Recorded By",
        ],
        [tuple(row) for row in rows],
    )

    # ---------------------------------------------------------
    # EEG AI
    # ---------------------------------------------------------

    rows = conn.execute(
        "SELECT p.patient_code, pr.created_at, "
        "pr.priority_label, pr.predicted_class, "
        "ROUND(pr.probability * 100, 1), u.user_id "
        "FROM predictions pr "
        "JOIN patients p ON pr.patient_id = p.id "
        "JOIN users u ON pr.created_by = u.id "
        f"{scope} "
        "ORDER BY pr.created_at DESC, pr.id DESC",
        params,
    ).fetchall()

    add_sheet(
        wb,
        "EEG_AI",
        [
            "Patient ID",
            "Time (UTC)",
            "AI-assisted priority indication",
            "Model class",
            "Model output probability (%)",
            "Analysed By",
        ],
        [tuple(row) for row in rows],
    )

    # ---------------------------------------------------------
    # Monitoring (SIMULATED)
    # ---------------------------------------------------------

    rows = conn.execute(
        "SELECT p.patient_code, m.created_at, m.scenario, "
        "m.step_index + 1, m.sim_source_class, m.priority_label, "
        "m.predicted_class, ROUND(m.probability * 100, 1), "
        "ROUND(m.mean_std, 3), u.user_id "
        "FROM monitoring_records m "
        "JOIN patients p ON m.patient_id = p.id "
        "JOIN users u ON m.created_by = u.id "
        f"{scope} "
        "ORDER BY p.patient_code, m.id",
        params,
    ).fetchall()

    add_sheet(
        wb,
        "Monitoring_SIMULATED",
        [
            "Patient ID",
            "Time (UTC)",
            "Scenario",
            "Step",
            "Simulated source class",
            "AI-assisted priority indication",
            "Model class",
            "Model output probability (%)",
            "Mean channel std",
            "Run by",
        ],
        [tuple(row) for row in rows],
    )

    # ---------------------------------------------------------
    # Notes
    # ---------------------------------------------------------

    notes = [
        (
            "AI-assisted results are decision-support indications "
            "from the NeuroTriage prototype.",
            "",
        ),
        (
            "Priority labels use a provisional, project-defined "
            "mapping from the EEG model classes.",
            "",
        ),
        (
            "Model output probabilities are not clinical confidence "
            "or a medical diagnosis.",
            "",
        ),
        (
            "The Monitoring_SIMULATED sheet contains simulated "
            "measurements, not data from a medical device.",
            "",
        ),
        (
            "Use synthetic/demo data only.",
            "",
        ),
    ]

    add_sheet(
        wb,
        "Notes",
        ["Note", "Value"],
        notes,
    )

    return wb


def save_workbook(workbook, user_id):
    """
    Save the generated workbook into backend/exports.
    """

    EXPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    safe_user_id = str(user_id).replace(
        "/", "_"
    ).replace(
        "\\", "_"
    )

    filename = (
        f"NeuroTriage_Export_"
        f"{safe_user_id}_"
        f"{timestamp}.xlsx"
    )

    path = EXPORT_DIR / filename

    workbook.save(path)

    return path